from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import classification_report, confusion_matrix

TRAINING_REPORT_WORKBOOK = "kaggle_model_training_report.xlsx"
SUMMARY_SHEET_NAME = "model_summary"
MODEL_SHEET_NAMES = {
    "XGBoost": "xgboost",
    "LightGBM": "lightgbm",
    "RandomForest": "randomforest",
}
MODEL_ORDER = ["XGBoost", "LightGBM", "RandomForest"]

LEGACY_TRAINING_CSV_FILES = [
    "xgboost_cv_fold_results.csv",
    "xgboost_cv_summary.csv",
    "xgboost_holdout_summary.csv",
    "xgboost_holdout_predictions.csv",
    "xgboost_confusion_matrix.csv",
    "xgboost_classification_report.csv",
    "lightgbm_cv_fold_results.csv",
    "lightgbm_cv_summary.csv",
    "lightgbm_holdout_summary.csv",
    "lightgbm_holdout_predictions.csv",
    "lightgbm_confusion_matrix.csv",
    "lightgbm_classification_report.csv",
    "randomforest_cv_fold_results.csv",
    "randomforest_cv_summary.csv",
    "randomforest_holdout_summary.csv",
    "randomforest_holdout_predictions.csv",
    "randomforest_confusion_matrix.csv",
    "randomforest_classification_report.csv",
]


def build_confusion_matrix_df(
    y_true: np.ndarray,
    y_pred: np.ndarray,
    class_names: np.ndarray,
) -> pd.DataFrame:
    labels_used = sorted(np.unique(np.concatenate([y_true, y_pred])).tolist())
    used_names = [class_names[i] for i in labels_used]
    matrix = confusion_matrix(y_true, y_pred, labels=labels_used)
    return pd.DataFrame(matrix, index=used_names, columns=used_names).reset_index().rename(
        columns={"index": "true_class"}
    )


def build_classification_report_df(
    y_true: np.ndarray,
    y_pred: np.ndarray,
    class_names: np.ndarray,
) -> pd.DataFrame:
    labels_used = sorted(np.unique(np.concatenate([y_true, y_pred])).tolist())
    used_names = [class_names[i] for i in labels_used]
    report = classification_report(
        y_true,
        y_pred,
        labels=labels_used,
        target_names=used_names,
        output_dict=True,
        zero_division=0,
    )
    rows = []
    for label_name, metrics in report.items():
        if isinstance(metrics, dict):
            rows.append({"label": label_name, **metrics})
    return pd.DataFrame(rows)


def build_summary_row(cv_summary_df: pd.DataFrame, holdout_df: pd.DataFrame) -> pd.DataFrame:
    holdout_renamed = holdout_df.rename(
        columns={
            "top1_accuracy": "holdout_top1_accuracy",
            "top3_accuracy": "holdout_top3_accuracy",
            "macro_f1": "holdout_macro_f1",
            "ece": "holdout_ece",
        }
    ).drop(columns=["split"], errors="ignore")
    summary_row = cv_summary_df.merge(holdout_renamed, on="model", how="left")
    ordered_columns = [
        "model",
        "cv_top1_mean",
        "cv_top1_std",
        "cv_top3_mean",
        "cv_top3_std",
        "cv_macro_f1_mean",
        "cv_macro_f1_std",
        "cv_ece_mean",
        "cv_ece_std",
        "cv_folds_used",
        "min_class_count_in_cv",
        "holdout_top1_accuracy",
        "holdout_top3_accuracy",
        "holdout_macro_f1",
        "holdout_ece",
    ]
    return summary_row[ordered_columns]


def load_existing_summary(workbook_path: Path) -> pd.DataFrame:
    if not workbook_path.exists():
        return pd.DataFrame()
    try:
        return pd.read_excel(workbook_path, sheet_name=SUMMARY_SHEET_NAME)
    except ValueError:
        return pd.DataFrame()


def merge_summary_rows(existing_summary: pd.DataFrame, new_summary_row: pd.DataFrame) -> pd.DataFrame:
    if existing_summary.empty:
        merged_summary = new_summary_row.copy()
    else:
        merged_summary = existing_summary.loc[existing_summary["model"] != new_summary_row.iloc[0]["model"]].copy()
        merged_summary = pd.concat([merged_summary, new_summary_row], ignore_index=True)

    merged_summary["model"] = pd.Categorical(merged_summary["model"], categories=MODEL_ORDER, ordered=True)
    merged_summary = merged_summary.sort_values("model", kind="stable").reset_index(drop=True)
    merged_summary["model"] = merged_summary["model"].astype(str)
    return merged_summary


def write_section(
    writer: pd.ExcelWriter,
    sheet_name: str,
    title: str,
    df: pd.DataFrame,
    startrow: int,
) -> int:
    pd.DataFrame({title: [title]}).to_excel(
        writer,
        sheet_name=sheet_name,
        startrow=startrow,
        index=False,
        header=False,
    )
    df.to_excel(writer, sheet_name=sheet_name, startrow=startrow + 1, index=False)
    return startrow + len(df) + 4


def remove_sheet_if_exists(workbook_path: Path, sheet_name: str) -> None:
    if not workbook_path.exists():
        return

    workbook = load_workbook(workbook_path)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        return

    worksheet = workbook[sheet_name]
    workbook.remove(worksheet)
    workbook.save(workbook_path)
    workbook.close()


def update_training_workbook(
    output_dir: Path,
    model_name: str,
    cv_df: pd.DataFrame,
    cv_summary_df: pd.DataFrame,
    holdout_df: pd.DataFrame,
    classification_report_df: pd.DataFrame,
    confusion_matrix_df: pd.DataFrame,
) -> Path:
    workbook_path = output_dir / TRAINING_REPORT_WORKBOOK
    detail_sheet_name = MODEL_SHEET_NAMES[model_name]

    existing_summary = load_existing_summary(workbook_path)
    new_summary_row = build_summary_row(cv_summary_df, holdout_df)
    merged_summary = merge_summary_rows(existing_summary, new_summary_row)

    writer_mode = "a" if workbook_path.exists() else "w"
    if writer_mode == "a":
        # In append mode, repeated writes to the same sheet with "replace"
        # would wipe earlier sections. Remove the target detail sheet once,
        # then use overlay so all sections land in a single fresh sheet.
        remove_sheet_if_exists(workbook_path, detail_sheet_name)
        writer_kwargs = {"engine": "openpyxl", "mode": "a", "if_sheet_exists": "overlay"}
    else:
        writer_kwargs = {"engine": "openpyxl", "mode": "w"}

    with pd.ExcelWriter(workbook_path, **writer_kwargs) as writer:
        merged_summary.to_excel(writer, sheet_name=SUMMARY_SHEET_NAME, index=False)

        current_row = 0
        current_row = write_section(writer, detail_sheet_name, "holdout_summary", holdout_df, current_row)
        current_row = write_section(writer, detail_sheet_name, "cv_fold_results", cv_df, current_row)
        current_row = write_section(
            writer,
            detail_sheet_name,
            "classification_report",
            classification_report_df,
            current_row,
        )
        write_section(writer, detail_sheet_name, "confusion_matrix", confusion_matrix_df, current_row)

    return workbook_path


def cleanup_legacy_training_csvs(output_dir: Path) -> None:
    for filename in LEGACY_TRAINING_CSV_FILES:
        path = output_dir / filename
        if path.exists():
            try:
                path.unlink()
            except PermissionError:
                print(f"[WARN] Could not remove locked legacy file: {path.name}")
