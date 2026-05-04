from __future__ import annotations

import ast
import json
import re
import time
from functools import lru_cache
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.preprocessing import LabelEncoder


BASE_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = BASE_DIR.parent
KAGGLE_DIR = PROJECT_ROOT / "1. Kaggle_Training"
ARTIFACTS_DIR = KAGGLE_DIR / "Artifacts_kaggle"
INPUT_DIR = BASE_DIR / "input" / "ddxplus"
OUTPUT_DIR = BASE_DIR / "output"

FEATURE_FILE = ARTIFACTS_DIR / "feature_names.pkl"
ENCODER_FILE = ARTIFACTS_DIR / "label_encoder.pkl"
CONDITIONS_FILE = INPUT_DIR / "release_conditions.json"
EVIDENCES_FILE = INPUT_DIR / "release_evidences.json"
SPLIT_FILES = {
    "train": INPUT_DIR / "release_train_patients.zip",
    "validate": INPUT_DIR / "release_validate_patients.zip",
    "test": INPUT_DIR / "release_test_patients.zip",
}

OUT_REPORT = OUTPUT_DIR / "ddxplus_alignment_report.xlsx"
OUT_REPORT_FALLBACK = OUTPUT_DIR / "ddxplus_alignment_report_latest.xlsx"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def timer(label: str, start: float) -> None:
    print(f"[TIMER] {label}: {time.perf_counter() - start:.2f} sec")


def normalize_text(text: str) -> str:
    text = str(text).strip().lower()
    text = text.replace("_", " ")
    text = text.replace("-", " ")
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def write_section(
    writer: pd.ExcelWriter,
    sheet_name: str,
    title: str,
    df: pd.DataFrame,
    startrow: int,
) -> int:
    pd.DataFrame({title: [title]}).to_excel(
        writer,
        sheet_name=sheet_name,
        startrow=startrow,
        index=False,
        header=False,
    )
    df.to_excel(writer, sheet_name=sheet_name, startrow=startrow + 1, index=False)
    return startrow + len(df) + 4


def save_alignment_workbook(
    workbook_path: Path,
    summary_df: pd.DataFrame,
    aligned_disease_counts_df: pd.DataFrame,
    disease_mapping_df: pd.DataFrame,
    feature_mapping_df: pd.DataFrame,
    alias_map_df: pd.DataFrame,
    split_coverage_df: pd.DataFrame,
    row_report_df: pd.DataFrame,
) -> Path:
    try:
        target_path = workbook_path
        with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
            current_row = 0
            current_row = write_section(writer, "summary", "alignment_summary", summary_df, current_row)
            current_row = write_section(writer, "summary", "aligned_disease_counts", aligned_disease_counts_df, current_row)
            write_section(writer, "summary", "split_coverage", split_coverage_df, current_row)

            disease_mapping_df.to_excel(writer, sheet_name="disease_mapping", index=False)
            feature_mapping_df.to_excel(writer, sheet_name="feature_mapping", index=False)
            alias_map_df.to_excel(writer, sheet_name="alias_map_reference", index=False)
            split_coverage_df.to_excel(writer, sheet_name="split_coverage", index=False)
            row_report_df.to_excel(writer, sheet_name="row_alignment_internal", index=False)
    except PermissionError:
        target_path = OUT_REPORT_FALLBACK
        with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
            current_row = 0
            current_row = write_section(writer, "summary", "alignment_summary", summary_df, current_row)
            current_row = write_section(writer, "summary", "aligned_disease_counts", aligned_disease_counts_df, current_row)
            write_section(writer, "summary", "split_coverage", split_coverage_df, current_row)

            disease_mapping_df.to_excel(writer, sheet_name="disease_mapping", index=False)
            feature_mapping_df.to_excel(writer, sheet_name="feature_mapping", index=False)
            alias_map_df.to_excel(writer, sheet_name="alias_map_reference", index=False)
            split_coverage_df.to_excel(writer, sheet_name="split_coverage", index=False)
            row_report_df.to_excel(writer, sheet_name="row_alignment_internal", index=False)

    workbook = load_workbook(target_path)
    if "row_alignment_internal" in workbook.sheetnames:
        workbook["row_alignment_internal"].sheet_state = "hidden"
        workbook.save(target_path)
    workbook.close()
    return target_path


DISEASE_ALIAS_SEED = {
    "Spontaneous pneumothorax": "pneumothorax",
    "GERD": "gastroesophageal reflux disease (gerd)",
    "Viral pharyngitis": "pharyngitis",
    "Guillain-Barré syndrome": "guillain barre syndrome",
    "Acute laryngitis": "laryngitis",
    "Allergic sinusitis": "chronic sinusitis",
    "Unstable angina": "angina",
    "Stable angina": "angina",
    "Bronchospasm / acute asthma exacerbation": "asthma",
    "Bronchitis": "acute bronchitis",
    "Acute COPD exacerbation / infection": "acute bronchitis",
    "URTI": "common cold",
    "Influenza": "flu",
    "Acute rhinosinusitis": "acute sinusitis",
    "Chronic rhinosinusitis": "chronic sinusitis",
    "Bronchiolitis": "acute bronchiolitis",
    "Pulmonary neoplasm": "lung cancer",
    "Possible NSTEMI / STEMI": "heart attack",
    "Pancreatic neoplasm": "pancreatic cancer",
    "Acute pulmonary edema": "pulmonary congestion",
    "Larygospasm": "acute bronchospasm",
}

BASE_SYMPTOM_MAP = {
    "E_91": ["fever"],
    "E_201": ["cough"],
    "E_77": ["coughing up sputum"],
    "E_66": ["shortness of breath"],
    "E_64": ["shortness of breath"],
    "E_214": ["wheezing"],
    "E_112": ["abnormal breathing sounds"],
    "E_194": ["abnormal breathing sounds"],
    "E_97": ["sore throat"],
    "E_181": ["nasal congestion"],
    "E_182": ["nasal congestion"],
    "E_94": ["chills"],
    "E_50": ["sweating"],
    "E_51": ["diarrhea"],
    "E_52": ["double vision"],
    "E_65": ["difficulty swallowing"],
    "E_63": ["difficulty speaking"],
    "E_144": ["muscle pain"],
    "E_161": ["decreased appetite"],
    "E_174": ["decreased appetite"],
    "E_175": ["fatigue"],
    "E_88": ["fatigue"],
    "E_89": ["fatigue"],
    "E_76": ["dizziness"],
    "E_82": ["dizziness"],
    "E_159": ["fainting"],
    "E_155": ["palpitations"],
    "E_172": ["abnormal movement of eyelid"],
    "E_9": ["swollen lymph nodes"],
    "E_212": ["hoarse voice"],
    "E_210": ["vomiting blood"],
    "E_206": ["mouth ulcer"],
    "E_148": ["nausea", "vomiting"],
    "E_129": ["skin rash"],
    "E_136": ["itching of skin"],
    "E_132": ["skin swelling"],
    "E_170": ["itchiness of eye"],
    "E_74": ["eye redness"],
    "E_45": ["coughing up sputum"],
    "E_30": ["abdominal distention"],
    "E_32": ["decreased appetite"],
    "E_173": ["heartburn"],
    "E_23": ["apnea"],
    "E_84": ["muscle weakness"],
    "E_176": ["weakness"],
    "E_157": ["focal weakness"],
    "E_177": ["focal weakness"],
    "E_192": ["neck stiffness or tightness"],
    "E_156": ["symptoms of the face"],
    "E_83": ["focal weakness"],
    "E_168": ["mouth pain"],
    "E_140": ["blood in stool"],
    "E_179": ["blood in stool"],
    "E_203": ["cough"],
    "E_202": ["cough"],
    "E_217": ["difficulty breathing"],
    "E_67": ["shortness of breath"],
    "E_75": ["difficulty breathing"],
    "E_128": ["difficulty breathing"],
}

PAIN_LOCATION_MAP = {
    "ear": "ear pain",
    "jaw": "mouth pain",
    "cheek": "facial pain",
    "forehead": "frontal headache",
    "temple": "headache",
    "back of head": "headache",
    "upper chest": "sharp chest pain",
    "lower chest": "sharp chest pain",
    "side of the chest": "sharp chest pain",
    "posterior chest wall": "sharp chest pain",
    "epigastric": "upper abdominal pain",
    "belly": "lower abdominal pain",
    "hypochondrium": "upper abdominal pain",
    "iliac fossa": "lower abdominal pain",
    "groin": "groin pain",
    "back of the neck": "neck pain",
    "side of the neck": "neck pain",
    "cervical spine": "neck pain",
    "thoracic spine": "back pain",
    "lumbar spine": "low back pain",
    "shoulder": "shoulder pain",
    "elbow": "elbow pain",
    "ankle": "ankle pain",
    "knee": "knee pain",
    "hip": "hip pain",
    "forearm": "arm pain",
    "hand": "hand or finger pain",
    "finger": "hand or finger pain",
    "foot": "foot or toe pain",
    "toe": "foot or toe pain",
    "calf": "leg pain",
    "thigh": "leg pain",
    "flank": "back pain",
    "scapula": "back pain",
    "mouth": "mouth pain",
    "pharynx": "sore throat",
    "tonsil": "sore throat",
}

SWELLING_LOCATION_MAP = {
    "ankle": "ankle swelling",
    "arm": "arm swelling",
    "elbow": "elbow swelling",
    "foot": "foot or toe swelling",
    "toe": "foot or toe swelling",
    "hand": "hand or finger swelling",
    "finger": "hand or finger swelling",
    "jaw": "jaw swelling",
    "knee": "knee swelling",
    "leg": "leg swelling",
    "neck": "neck swelling",
    "shoulder": "shoulder swelling",
    "throat": "throat swelling",
    "wrist": "hand or finger swelling",
    "eye": "swollen eye",
    "eyelid": "eyelid swelling",
    "skin": "skin swelling",
}


def main() -> None:
    print("--- Starting DDXPlus Alignment ---")
    script_start = time.perf_counter()

    try:
        section_start = time.perf_counter()
        try:
            feature_names: list[str] = joblib.load(FEATURE_FILE)
            encoder: LabelEncoder = joblib.load(ENCODER_FILE)
            conditions = json.loads(CONDITIONS_FILE.read_text(encoding="utf-8"))
            evidences = json.loads(EVIDENCES_FILE.read_text(encoding="utf-8"))
            raw_splits = {
                split_name: pd.read_csv(path, compression="zip")
                for split_name, path in SPLIT_FILES.items()
            }
        finally:
            timer("load_resources", section_start)

        feature_set = set(feature_names)
        valid_disease_lookup = {normalize_text(name): name for name in encoder.classes_.tolist()}
        disease_alias_map = {
            normalize_text(raw_alias): canonical
            for raw_alias, canonical in DISEASE_ALIAS_SEED.items()
            if canonical in encoder.classes_
        }
        base_symptom_map = {
            code: [feature for feature in mapped_features if feature in feature_set]
            for code, mapped_features in BASE_SYMPTOM_MAP.items()
        }
        pain_value_lookup = {key: value["en"].lower() for key, value in evidences["E_55"]["value_meaning"].items()}
        swelling_value_lookup = {key: value["en"].lower() for key, value in evidences["E_152"]["value_meaning"].items()}

        alias_map_rows = []
        for raw_alias, mapped_name in sorted(disease_alias_map.items()):
            alias_map_rows.append(
                {
                    "map_scope": "disease",
                    "raw_value": raw_alias,
                    "mapped_value": mapped_name,
                }
            )
        alias_map_rows.extend(
            [
                {
                    "map_scope": "symptom_group",
                    "raw_value": "E_55",
                    "mapped_value": "location dependent pain mapping",
                },
                {
                    "map_scope": "symptom_group",
                    "raw_value": "E_152",
                    "mapped_value": "location dependent swelling mapping",
                },
            ]
        )
        alias_map_df = pd.DataFrame(alias_map_rows)

        def map_disease(raw_name: str) -> tuple[str | None, str]:
            normalized = normalize_text(raw_name)
            if normalized in valid_disease_lookup:
                return valid_disease_lookup[normalized], "exact_match"
            if normalized in disease_alias_map:
                return disease_alias_map[normalized], "alias_map"
            return None, "unmapped"

        def map_pain_location(value_code: str) -> str | None:
            location_text = pain_value_lookup.get(value_code, "")
            for needle, feature in PAIN_LOCATION_MAP.items():
                if needle in location_text and feature in feature_set:
                    return feature
            return None

        def map_swelling_location(value_code: str) -> str | None:
            location_text = swelling_value_lookup.get(value_code, "")
            for needle, feature in SWELLING_LOCATION_MAP.items():
                if needle in location_text and feature in feature_set:
                    return feature
            return None

        @lru_cache(maxsize=300_000)
        def map_evidence_text(evidence_text: str) -> tuple[str, int]:
            features_used: set[str] = set()
            evidence_items = ast.literal_eval(evidence_text)
            for item in evidence_items:
                if "_@_" in item:
                    code, value = item.split("_@_", 1)
                else:
                    code, value = item, None

                root_code = evidences.get(code, {}).get("code_question", code)
                for feature in base_symptom_map.get(root_code, []):
                    features_used.add(feature)

                if code == "E_55" and value:
                    mapped_feature = map_pain_location(value)
                    if mapped_feature:
                        features_used.add(mapped_feature)

                if code == "E_152" and value:
                    mapped_feature = map_swelling_location(value)
                    if mapped_feature:
                        features_used.add(mapped_feature)

            ordered = sorted(features_used)
            return " | ".join(ordered), len(ordered)

        section_start = time.perf_counter()
        try:
            combined_raw_df = pd.concat(
                [
                    split_df.assign(source_split=split_name).rename(
                        columns={"PATHOLOGY": "raw_pathology", "EVIDENCES": "evidences_text"}
                    )
                    for split_name, split_df in raw_splits.items()
                ],
                ignore_index=True,
            )

            raw_condition_counts = combined_raw_df["raw_pathology"].value_counts().to_dict()
            disease_mapping_rows = []
            for raw_condition_name in conditions.keys():
                mapped_disease, mapping_status = map_disease(raw_condition_name)
                disease_mapping_rows.append(
                    {
                        "raw_pathology": raw_condition_name,
                        "mapped_disease": mapped_disease,
                        "mapping_status": mapping_status,
                        "row_count": int(raw_condition_counts.get(raw_condition_name, 0)),
                    }
                )
            disease_mapping_df = pd.DataFrame(disease_mapping_rows)

            feature_mapping_rows = []
            symptom_groups = sorted(
                {
                    evidence["code_question"]
                    for evidence in evidences.values()
                    if not evidence["is_antecedent"]
                }
            )
            for code_question in symptom_groups:
                evidence = evidences[code_question]
                mapped_features = base_symptom_map.get(code_question, [])
                if code_question == "E_55":
                    mapping_notes = "location dependent pain mapping"
                elif code_question == "E_152":
                    mapping_notes = "location dependent swelling mapping"
                else:
                    mapping_notes = ""
                feature_mapping_rows.append(
                    {
                        "ddxplus_evidence_code": code_question,
                        "question_en": evidence["question_en"],
                        "mapped_features": " | ".join(mapped_features),
                        "mapping_notes": mapping_notes,
                        "mapping_status": "mapped" if mapped_features or mapping_notes else "unmapped",
                    }
                )
            feature_mapping_df = pd.DataFrame(feature_mapping_rows)
        finally:
            timer("build_mapping_tables", section_start)

        section_start = time.perf_counter()
        try:
            row_report_frames = []
            split_stats_rows = []
            for split_name, split_df in raw_splits.items():
                working_df = split_df.rename(columns={"PATHOLOGY": "raw_pathology", "EVIDENCES": "evidences_text"}).copy()
                working_df["source_split"] = split_name
                working_df["row_id"] = np.arange(len(working_df))

                mapped_pairs = working_df["raw_pathology"].map(map_disease)
                working_df["mapped_disease"] = mapped_pairs.map(lambda item: item[0])
                working_df["disease_mapping_status"] = mapped_pairs.map(lambda item: item[1])
                mapped_features = working_df["evidences_text"].map(map_evidence_text)
                working_df["mapped_features"] = mapped_features.map(lambda item: item[0])
                working_df["mapped_feature_count"] = mapped_features.map(lambda item: item[1]).astype(int)
                working_df["retained_for_alignment"] = (
                    working_df["mapped_disease"].notna() & working_df["mapped_feature_count"].gt(0)
                ).astype(int)

                retained_df = working_df.loc[working_df["retained_for_alignment"] == 1].copy()
                dedup_df = retained_df.drop_duplicates(subset=["mapped_disease", "mapped_features"], keep="first").copy()
                dedup_index_set = set(dedup_df.index.tolist())
                working_df["retained_after_dedup_within_split"] = working_df.index.map(
                    lambda idx: int(idx in dedup_index_set)
                )

                row_report_frames.append(
                    working_df[
                        [
                            "source_split",
                            "row_id",
                            "raw_pathology",
                            "mapped_disease",
                            "disease_mapping_status",
                            "mapped_feature_count",
                            "retained_for_alignment",
                            "retained_after_dedup_within_split",
                            "mapped_features",
                        ]
                    ].copy()
                )
                split_stats_rows.append(
                    {
                        "source_split": split_name,
                        "raw_rows": len(working_df),
                        "mapped_rows_before_dedup": len(retained_df),
                        "rows_after_dedup_within_split": len(dedup_df),
                        "raw_classes": int(working_df["raw_pathology"].nunique()),
                        "mapped_classes_after_dedup": int(dedup_df["mapped_disease"].nunique()),
                    }
                )

            row_report_df = pd.concat(row_report_frames, ignore_index=True)
            split_coverage_df = pd.DataFrame(split_stats_rows)
        finally:
            timer("row_alignment", section_start)

        section_start = time.perf_counter()
        try:
            dedup_retained_df = row_report_df.loc[row_report_df["retained_after_dedup_within_split"] == 1].copy()
            aligned_disease_counts_df = (
                dedup_retained_df["mapped_disease"]
                .value_counts()
                .rename_axis("mapped_disease")
                .reset_index(name="aligned_row_count")
            )
            summary_df = pd.DataFrame(
                [
                    {"metric": "raw_rows_total", "value": len(row_report_df)},
                    {"metric": "raw_disease_classes", "value": len(conditions)},
                    {"metric": "mapped_disease_classes_after_split_dedup", "value": int(dedup_retained_df["mapped_disease"].nunique())},
                    {"metric": "aligned_feature_space_size", "value": len(feature_names)},
                    {"metric": "aligned_label_space_size", "value": len(encoder.classes_)},
                    {"metric": "mapped_symptom_groups", "value": int(feature_mapping_df.loc[feature_mapping_df["mapping_status"] == "mapped"].shape[0])},
                    {"metric": "unmapped_symptom_groups", "value": int(feature_mapping_df.loc[feature_mapping_df["mapping_status"] == "unmapped"].shape[0])},
                ]
            )

            row_report_export_df = row_report_df.loc[
                row_report_df["retained_after_dedup_within_split"] == 1
            ].copy().reset_index(drop=True)

            target_path = save_alignment_workbook(
                workbook_path=OUT_REPORT,
                summary_df=summary_df,
                aligned_disease_counts_df=aligned_disease_counts_df,
                disease_mapping_df=disease_mapping_df,
                feature_mapping_df=feature_mapping_df,
                alias_map_df=alias_map_df,
                split_coverage_df=split_coverage_df,
                row_report_df=row_report_export_df,
            )
            print(f"Workbook saved: {target_path.name}")
        finally:
            timer("save_outputs", section_start)
    finally:
        timer("total_script", script_start)


if __name__ == "__main__":
    main()
