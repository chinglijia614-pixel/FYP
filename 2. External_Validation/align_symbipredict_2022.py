from __future__ import annotations

import re
import time
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from sklearn.preprocessing import LabelEncoder
from openpyxl import load_workbook

# =============================================================================
# CONFIG
# =============================================================================
BASE_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = BASE_DIR.parent
KAGGLE_DIR = PROJECT_ROOT / "1. Kaggle_Training"
ARTIFACTS_DIR = KAGGLE_DIR / "Artifacts_kaggle"
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"

RAW_FILE = INPUT_DIR / "symbipredict_2022.csv"
FEATURE_FILE = ARTIFACTS_DIR / "feature_names.pkl"
ENCODER_FILE = ARTIFACTS_DIR / "label_encoder.pkl"
TRAIN_X_FILE = KAGGLE_DIR / "input" / "X_train.csv"
TRAIN_Y_FILE = KAGGLE_DIR / "input" / "y_train.csv"
TEST_X_FILE = KAGGLE_DIR / "input" / "X_test.csv"
TEST_Y_FILE = KAGGLE_DIR / "input" / "y_test.csv"

OUT_REPORT = OUTPUT_DIR / "symbipredict_alignment_report.xlsx"
OUT_REPORT_FALLBACK = OUTPUT_DIR / "symbipredict_alignment_report_latest.xlsx"

for d in [OUTPUT_DIR]:
    d.mkdir(parents=True, exist_ok=True)


# =============================================================================
# HELPERS
# =============================================================================
def timer(label: str, start: float) -> None:
    print(f"[TIMER] {label}: {time.perf_counter() - start:.2f} sec")


def normalize_text(text: str) -> str:
    text = str(text).strip().lower()
    text = text.replace("_", " ")
    text = text.replace("-", " ")
    text = text.replace("&", " and ")
    text = re.sub(r"[()/]", " ", text)
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def write_section(
    writer: pd.ExcelWriter,
    sheet_name: str,
    title: str,
    df: pd.DataFrame,
    startrow: int,
) -> int:
    pd.DataFrame({title: [title]}).to_excel(
        writer,
        sheet_name=sheet_name,
        startrow=startrow,
        index=False,
        header=False,
    )
    df.to_excel(writer, sheet_name=sheet_name, startrow=startrow + 1, index=False)
    return startrow + len(df) + 4


def resolve_alias_targets(alias_seed: dict[str, str], lookup: dict[str, str]) -> dict[str, str]:
    resolved = {}
    for raw_alias, canonical_target in alias_seed.items():
        alias_key = normalize_text(raw_alias)
        target_key = normalize_text(canonical_target)
        if target_key in lookup:
            resolved[alias_key] = lookup[target_key]
    return resolved


def to_mapping_status(mapping_type: str) -> str:
    status_lookup = {
        "exact": "exact_match",
        "alias": "alias_map",
        "unmapped": "unmapped",
    }
    return status_lookup.get(mapping_type, mapping_type)


def save_alignment_workbook(
    workbook_path: Path,
    summary_df: pd.DataFrame,
    aligned_disease_counts_df: pd.DataFrame,
    disease_mapping_df: pd.DataFrame,
    feature_mapping_df: pd.DataFrame,
    alias_map_df: pd.DataFrame,
    symptom_count_comparison_df: pd.DataFrame,
    row_report_df: pd.DataFrame,
) -> Path:
    try:
        target_path = workbook_path
        with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
            current_row = 0
            current_row = write_section(writer, "summary", "alignment_summary", summary_df, current_row)
            write_section(writer, "summary", "aligned_disease_counts", aligned_disease_counts_df, current_row)

            disease_mapping_df.to_excel(writer, sheet_name="disease_mapping", index=False)
            feature_mapping_df.to_excel(writer, sheet_name="symptom_mapping", index=False)
            alias_map_df.to_excel(writer, sheet_name="alias_map_reference", index=False)
            symptom_count_comparison_df.to_excel(
                writer,
                sheet_name="mapped_disease_symptom_counts",
                index=False,
            )
            row_report_df.to_excel(writer, sheet_name="row_alignment_internal", index=False)
    except PermissionError:
        target_path = OUT_REPORT_FALLBACK
        with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
            current_row = 0
            current_row = write_section(writer, "summary", "alignment_summary", summary_df, current_row)
            write_section(writer, "summary", "aligned_disease_counts", aligned_disease_counts_df, current_row)

            disease_mapping_df.to_excel(writer, sheet_name="disease_mapping", index=False)
            feature_mapping_df.to_excel(writer, sheet_name="symptom_mapping", index=False)
            alias_map_df.to_excel(writer, sheet_name="alias_map_reference", index=False)
            symptom_count_comparison_df.to_excel(
                writer,
                sheet_name="mapped_disease_symptom_counts",
                index=False,
            )
            row_report_df.to_excel(writer, sheet_name="row_alignment_internal", index=False)

    workbook = load_workbook(target_path)
    if "row_alignment_internal" in workbook.sheetnames:
        workbook["row_alignment_internal"].sheet_state = "hidden"
        workbook.save(target_path)
    workbook.close()
    return target_path


def build_visible_disease_mapping_df(disease_mapping_df: pd.DataFrame) -> pd.DataFrame:
    return disease_mapping_df[
        ["raw_prognosis", "mapped_disease", "mapping_status", "row_count"]
    ].copy()


def build_visible_symptom_mapping_df(feature_mapping_df: pd.DataFrame) -> pd.DataFrame:
    return feature_mapping_df[
        [
            "raw_feature",
            "mapped_feature",
            "mapping_status",
            "active_row_count_mapped_diseases",
        ]
    ].copy()


def build_visible_symptom_count_comparison_df(symptom_count_comparison_df: pd.DataFrame) -> pd.DataFrame:
    return symptom_count_comparison_df[
        [
            "raw_prognosis",
            "mapped_disease",
            "mapping_status",
            "symbipredict_unique_raw_symptom_count",
            "symbipredict_unique_mapped_symptom_count",
            "kaggle_rows",
            "kaggle_unique_symptom_count",
            "mapped_symptom_gap_vs_kaggle",
        ]
    ].copy()


# =============================================================================
# CONSERVATIVE ALIAS MAPS
# =============================================================================
DISEASE_ALIAS_SEED = {
    # Keep only near-equivalent naming variants.
    "Bronchial Asthma": "asthma",
    "Dimorphic Hemmorhoids (piles)": "hemorrhoids",
    "GERD": "gastroesophageal reflux disease (gerd)",
}

SYMPTOM_ALIAS_SEED = {
    # Keep symptom aliases strict as well: only near-identical wording variants.
    "itching": "itching of skin",
    "continuous_sneezing": "sneezing",
    "burning_micturition": "painful urination",
    "breathlessness": "shortness of breath",
    "loss_of_appetite": "decreased appetite",
    "swelled_lymph_nodes": "swollen lymph nodes",
    "excessive_hunger": "excessive appetite",
    "diarrhoea": "diarrhea",
    "fast_heart_rate": "increased heart rate",
    "slurred_speech": "difficulty speaking",
}


# =============================================================================
# MAPPING
# =============================================================================
def map_disease(
    raw_name: str,
    valid_disease_lookup: dict[str, str],
    disease_alias_map: dict[str, str],
) -> tuple[str | None, str]:
    key = normalize_text(raw_name)
    if not key:
        return None, "unmapped"

    if key in valid_disease_lookup:
        return valid_disease_lookup[key], "exact"

    if key in disease_alias_map:
        return disease_alias_map[key], "alias"

    return None, "unmapped"


def map_symptom_column(
    raw_column: str,
    feature_lookup: dict[str, str],
    symptom_alias_map: dict[str, str],
) -> tuple[str | None, str]:
    key = normalize_text(raw_column)
    if not key:
        return None, "unmapped"

    if key in feature_lookup:
        return feature_lookup[key], "exact"

    if key in symptom_alias_map:
        return symptom_alias_map[key], "alias"

    return None, "unmapped"


# =============================================================================
# MAIN
# =============================================================================
def main() -> None:
    print("--- Starting SymbiPredict 2022 Alignment ---")
    script_start = time.perf_counter()

    try:
        section_start = time.perf_counter()
        try:
            df_raw = pd.read_csv(RAW_FILE)
            model_features: list[str] = joblib.load(FEATURE_FILE)
            encoder: LabelEncoder = joblib.load(ENCODER_FILE)
            X_train = pd.read_csv(TRAIN_X_FILE)
            y_train = pd.read_csv(TRAIN_Y_FILE)["target"]
            X_test = pd.read_csv(TEST_X_FILE)
            y_test = pd.read_csv(TEST_Y_FILE)["target"]
        finally:
            timer("load_resources", section_start)

        valid_diseases = list(encoder.classes_)
        valid_disease_lookup = {normalize_text(x): x for x in valid_diseases}
        feature_lookup = {normalize_text(x): x for x in model_features}
        disease_alias_map = resolve_alias_targets(DISEASE_ALIAS_SEED, valid_disease_lookup)
        symptom_alias_map = resolve_alias_targets(SYMPTOM_ALIAS_SEED, feature_lookup)

        alias_map_rows = []
        for raw_alias, mapped_name in sorted(disease_alias_map.items()):
            alias_map_rows.append(
                {
                    "map_scope": "disease",
                    "raw_value": raw_alias,
                    "mapped_value": mapped_name,
                }
            )
        for raw_alias, mapped_name in sorted(symptom_alias_map.items()):
            alias_map_rows.append(
                {
                    "map_scope": "symptom",
                    "raw_value": raw_alias,
                    "mapped_value": mapped_name,
                }
            )
        alias_map_df = pd.DataFrame(alias_map_rows)

        raw_feature_cols = [col for col in df_raw.columns if col != "prognosis"]

        section_start = time.perf_counter()
        try:
            disease_mapping_rows = []
            prognosis_values = sorted(df_raw["prognosis"].dropna().astype(str).unique().tolist())
            exact_disease_labels = 0
            alias_disease_labels = 0

            for prognosis in prognosis_values:
                mapped_disease, mapping_type = map_disease(
                    prognosis,
                    valid_disease_lookup,
                    disease_alias_map,
                )
                row_count = int((df_raw["prognosis"].astype(str) == prognosis).sum())

                if mapping_type == "exact":
                    exact_disease_labels += 1
                elif mapping_type == "alias":
                    alias_disease_labels += 1

                disease_mapping_rows.append(
                    {
                        "raw_prognosis": prognosis,
                        "mapped_disease": mapped_disease,
                        "mapping_type": mapping_type,
                        "mapping_status": to_mapping_status(mapping_type),
                        "row_count": row_count,
                    }
                )

            disease_mapping_df = pd.DataFrame(disease_mapping_rows)

            mapped_raw_prognosis_values = set(
                disease_mapping_df.loc[disease_mapping_df["mapped_disease"].notna(), "raw_prognosis"].astype(str).tolist()
            )
            mapped_disease_mask = df_raw["prognosis"].astype(str).isin(mapped_raw_prognosis_values)

            feature_mapping_rows = []
            feature_map: dict[str, str] = {}
            exact_feature_cols = 0
            alias_feature_cols = 0

            for col in raw_feature_cols:
                mapped_feature, mapping_type = map_symptom_column(col, feature_lookup, symptom_alias_map)
                active_values = pd.to_numeric(df_raw[col], errors="coerce").fillna(0).astype(int)
                active_row_count = int(active_values.gt(0).sum())
                active_row_count_mapped_diseases = int(active_values.loc[mapped_disease_mask].gt(0).sum())

                if mapped_feature is not None:
                    feature_map[col] = mapped_feature
                    if mapping_type == "exact":
                        exact_feature_cols += 1
                    elif mapping_type == "alias":
                        alias_feature_cols += 1

                feature_mapping_rows.append(
                    {
                        "raw_feature": col,
                        "mapped_feature": mapped_feature,
                        "mapping_type": mapping_type,
                        "mapping_status": to_mapping_status(mapping_type),
                        "active_row_count": active_row_count,
                        "active_row_count_mapped_diseases": active_row_count_mapped_diseases,
                        "active_in_mapped_diseases": int(active_row_count_mapped_diseases > 0),
                    }
                )

            feature_mapping_df = pd.DataFrame(feature_mapping_rows)
        finally:
            timer("build_mapping_tables", section_start)

        section_start = time.perf_counter()
        try:
            row_report_rows = []

            for idx, row in df_raw.iterrows():
                raw_prognosis = row["prognosis"]
                mapped_disease, disease_mapping_type = map_disease(
                    raw_prognosis,
                    valid_disease_lookup,
                    disease_alias_map,
                )

                active_raw_features = []
                mapped_features = []
                unmapped_active_features = []

                for raw_col in raw_feature_cols:
                    value = pd.to_numeric(row[raw_col], errors="coerce")
                    if pd.notna(value) and int(value) > 0:
                        active_raw_features.append(raw_col)
                        if raw_col in feature_map:
                            mapped_features.append(feature_map[raw_col])
                        else:
                            unmapped_active_features.append(raw_col)

                mapped_features = sorted(set(mapped_features))

                retained = mapped_disease is not None and len(active_raw_features) > 0

                row_report_rows.append(
                    {
                        "row_id": idx,
                        "raw_prognosis": raw_prognosis,
                        "mapped_disease": mapped_disease,
                        "disease_mapping_type": disease_mapping_type,
                        "disease_mapping_status": to_mapping_status(disease_mapping_type),
                        "active_raw_feature_count": len(active_raw_features),
                        "mapped_feature_count": len(mapped_features),
                        "retained_for_alignment": int(retained),
                        "active_raw_features": " | ".join(active_raw_features),
                        "mapped_features": " | ".join(mapped_features),
                        "unmapped_active_features": " | ".join(unmapped_active_features),
                    }
                )

            row_report_df = pd.DataFrame(row_report_rows)
            retained_df = row_report_df.loc[row_report_df["retained_for_alignment"] == 1].copy().reset_index(drop=True)
        finally:
            timer("row_alignment", section_start)

        section_start = time.perf_counter()
        try:
            unmapped_active_feature_counts: dict[str, int] = {}
            for raw_text in retained_df["unmapped_active_features"]:
                items = [x.strip() for x in str(raw_text).split("|") if x.strip()]
                for item in items:
                    unmapped_active_feature_counts[item] = unmapped_active_feature_counts.get(item, 0) + 1

            aligned_disease_counts_df = (
                retained_df["mapped_disease"]
                .value_counts()
                .rename_axis("mapped_disease")
                .reset_index(name="aligned_row_count")
            )

            kaggle_full_X = pd.concat([X_train, X_test], axis=0, ignore_index=True)
            kaggle_full_y = pd.concat([y_train, y_test], axis=0, ignore_index=True)
            kaggle_disease_labels = encoder.inverse_transform(kaggle_full_y.astype(int).to_numpy())

            kaggle_profile_df = kaggle_full_X.copy()
            kaggle_profile_df["mapped_disease"] = kaggle_disease_labels

            kaggle_symptom_counts = []
            for mapped_disease in sorted(disease_mapping_df["mapped_disease"].dropna().unique().tolist()):
                disease_rows = kaggle_profile_df.loc[kaggle_profile_df["mapped_disease"] == mapped_disease, model_features]
                kaggle_symptom_counts.append(
                    {
                        "mapped_disease": mapped_disease,
                        "kaggle_rows": int(len(disease_rows)),
                        "kaggle_unique_symptom_count": int(disease_rows.any(axis=0).sum()),
                    }
                )
            kaggle_symptom_comparison_df = pd.DataFrame(kaggle_symptom_counts)

            mapped_disease_symptom_rows = []
            for _, mapping_row in disease_mapping_df.loc[disease_mapping_df["mapped_disease"].notna()].iterrows():
                raw_name = mapping_row["raw_prognosis"]
                mapped_disease = mapping_row["mapped_disease"]
                disease_row_subset = row_report_df.loc[row_report_df["raw_prognosis"] == raw_name].copy()

                raw_feature_set = set()
                mapped_feature_set = set()
                retained_external_feature_set = set()

                for raw_text in disease_row_subset["active_raw_features"]:
                    raw_feature_set.update(x.strip() for x in str(raw_text).split("|") if x.strip())

                for mapped_text in disease_row_subset["mapped_features"]:
                    mapped_feature_set.update(x.strip() for x in str(mapped_text).split("|") if x.strip())

                for raw_text in disease_row_subset["unmapped_active_features"]:
                    retained_external_feature_set.update(x.strip() for x in str(raw_text).split("|") if x.strip())

                mapped_disease_symptom_rows.append(
                    {
                        "raw_prognosis": raw_name,
                        "mapped_disease": mapped_disease,
                        "mapping_type": mapping_row["mapping_type"],
                        "mapping_status": mapping_row["mapping_status"],
                        "symbipredict_rows": int(len(disease_row_subset)),
                        "symbipredict_retained_rows": int(disease_row_subset["retained_for_alignment"].sum()),
                        "symbipredict_unique_raw_symptom_count": int(len(raw_feature_set)),
                        "symbipredict_unique_mapped_symptom_count": int(len(mapped_feature_set)),
                        "symbipredict_unique_retained_external_symptom_count": int(len(retained_external_feature_set)),
                        "symbipredict_union_symptom_count": int(len(mapped_feature_set) + len(retained_external_feature_set)),
                    }
                )

            symptom_count_comparison_df = pd.DataFrame(mapped_disease_symptom_rows).merge(
                kaggle_symptom_comparison_df,
                on="mapped_disease",
                how="left",
            )
            symptom_count_comparison_df["mapped_symptom_gap_vs_kaggle"] = (
                symptom_count_comparison_df["symbipredict_unique_mapped_symptom_count"]
                - symptom_count_comparison_df["kaggle_unique_symptom_count"]
            )
            symptom_count_comparison_df["mapped_symptom_ratio_vs_kaggle"] = (
                symptom_count_comparison_df["symbipredict_unique_mapped_symptom_count"]
                / symptom_count_comparison_df["kaggle_unique_symptom_count"]
            ).replace([np.inf, -np.inf], np.nan)

            visible_disease_mapping_df = build_visible_disease_mapping_df(disease_mapping_df)
            visible_feature_mapping_df = build_visible_symptom_mapping_df(feature_mapping_df)
            visible_symptom_count_comparison_df = build_visible_symptom_count_comparison_df(
                symptom_count_comparison_df
            )

            summary_df = pd.DataFrame(
                [
                    {"metric": "raw_rows", "value": len(df_raw)},
                    {"metric": "aligned_rows", "value": len(retained_df)},
                    {"metric": "dropped_rows", "value": len(df_raw) - len(retained_df)},
                    {"metric": "raw_feature_columns", "value": len(raw_feature_cols)},
                    {"metric": "mapped_feature_columns_exact", "value": exact_feature_cols},
                    {"metric": "mapped_feature_columns_alias", "value": alias_feature_cols},
                    {"metric": "mapped_feature_columns_total", "value": len(feature_map)},
                    {
                        "metric": "raw_feature_columns_active_in_mapped_diseases",
                        "value": int(feature_mapping_df["active_in_mapped_diseases"].sum()),
                    },
                    {
                        "metric": "retained_external_only_feature_columns",
                        "value": int(len(unmapped_active_feature_counts)),
                    },
                    {"metric": "raw_prognosis_labels", "value": len(prognosis_values)},
                    {"metric": "mapped_prognosis_labels_exact", "value": exact_disease_labels},
                    {"metric": "mapped_prognosis_labels_alias", "value": alias_disease_labels},
                    {
                        "metric": "mapped_prognosis_labels_total",
                        "value": int(disease_mapping_df["mapped_disease"].notna().sum()),
                    },
                    {"metric": "aligned_feature_space_size", "value": len(model_features)},
                    {"metric": "aligned_label_space_size", "value": len(valid_diseases)},
                ]
            )

            saved_report_path = save_alignment_workbook(
                workbook_path=OUT_REPORT,
                summary_df=summary_df,
                aligned_disease_counts_df=aligned_disease_counts_df,
                disease_mapping_df=visible_disease_mapping_df,
                feature_mapping_df=visible_feature_mapping_df,
                alias_map_df=alias_map_df,
                symptom_count_comparison_df=visible_symptom_count_comparison_df,
                row_report_df=row_report_df,
            )
        finally:
            timer("save_outputs", section_start)

        print(f"Raw rows: {len(df_raw)}")
        print(f"Aligned rows kept: {len(retained_df)}")
        print(
            "Mapped diseases: "
            f"{int(disease_mapping_df['mapped_disease'].notna().sum())} / {len(prognosis_values)}"
        )
        print(
            "Mapped symptoms in mapped diseases: "
            f"{int(feature_mapping_df.loc[feature_mapping_df['active_in_mapped_diseases'] == 1, 'mapped_feature'].notna().sum())}"
            f" / {int(feature_mapping_df['active_in_mapped_diseases'].sum())}"
        )
        print(
            "Retained external-only symptoms in mapped diseases: "
            f"{int(len(retained_df['unmapped_active_features'].astype(str).str.split('|').explode().str.strip().replace('', np.nan).dropna().unique()))}"
        )
        print("Saved:")
        print(f"  - {saved_report_path.name}")
    finally:
        timer("total_script", script_start)


if __name__ == "__main__":
    main()
